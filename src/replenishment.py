"""Replenishment logic for the footwear stock chatbot.

Loads and cleans the stock dataset, and exposes the query functions used by the
chatbot. Every query function takes the cleaned dataframe as its first argument,
so the module holds no hidden state and can be safely cached by Streamlit.
"""

import re
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Constants
# --------------------------------------------------------------------------

# project root, resolved from this file so paths work from any working directory
PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_PATH = PROJECT_ROOT / 'data' / 'stock_snapshot.xlsx'
OUTPUT_DIR = PROJECT_ROOT / 'outputs'

# the 6 stores selected by volume, plus HQ as the stock source
SELECTED_STORES = ['5-52', '5-53', '5-50', '2-8', '5-51', '2-41']
HQ_ID = '1-HQ'

STORE_NAMES = {
    '5-52': 'Outlet Cascais',
    '5-53': 'Outlet Gaia',
    '5-50': 'Outlet Loures',
    '2-8': 'Loja Amoreiras',
    '5-51': 'Outlet Almada',
    '2-41': 'Loja Colombo',
    '1-HQ': 'HQ',
}

# warehouse codes: 1 = central warehouse, 2 = stores, 5 = outlets
WAREHOUSE_TYPE_LABELS = {1: 'hq', 2: 'store', 5: 'outlet'}
STORE_TYPE_DISPLAY = {'hq': 'HQ', 'store': 'Store', 'outlet': 'Outlet'}

USEFUL_COLUMNS = ['MARCA', 'ARM', 'REF', 'DESCRICAO', 'TAMANHO',
                  'LOCALIZACAO', 'QUANTIDADE', 'TIPO ARTIGO', 'FAMILIA']

COLUMN_RENAMES = {
    'MARCA': 'BRAND',
    'ARM': 'WAREHOUSE',
    'REF': 'REFERENCE',
    'DESCRICAO': 'DESCRIPTION',
    'TAMANHO': 'SIZE',
    'LOCALIZACAO': 'LOCATION',
    'QUANTIDADE': 'QUANTITY',
    'TIPO ARTIGO': 'ITEM_TYPE',
    'FAMILIA': 'FAMILY',
}

BRAND_FIXES = {
    'Fila apparel_acc': 'Fila',
    'Merrell Aces': 'Merrell Foot',
    'Munich Sports': 'Munich',
}


# --------------------------------------------------------------------------
# Size conversion
# --------------------------------------------------------------------------

def convert_size_final(raw):
    """Convert every SIZE format found in the source data into one numeric scale.

    Kids Y/C codes (US sizing) keep their number, which is not equivalent to EU
    adult sizing but is a documented simplification. Letter sizes (S, M, UNI)
    have no numeric equivalent and return NaN.
    """
    s = str(raw).strip()
    if re.fullmatch(r'\d+,\d+', s):
        return float(s.replace(',', '.'))
    if re.fullmatch(r'\d+\.\d+', s):
        return float(s)
    if re.fullmatch(r'\d+-\d+', s):
        a, b = s.split('-')
        return (float(a) + float(b)) / 2
    if re.fullmatch(r'\d{4}', s):
        a, b = s[:2], s[2:]
        return (float(a) + float(b)) / 2
    if re.fullmatch(r'\d{3}', s):
        return int(s) / 10
    if re.fullmatch(r'\d{1,2}', s):
        return float(s)
    if re.fullmatch(r'\d+(\.\d+)?[YC]', s):
        return float(s[:-1])
    return np.nan


def format_size_label(size):
    """Format a size number as clean text: whole numbers drop the decimal."""
    return str(int(size)) if float(size).is_integer() else str(size)


# --------------------------------------------------------------------------
# Data loading
# --------------------------------------------------------------------------

def load_data(path=None):
    """Load the raw dataset and apply every cleaning step decided in the EDA.

    Mirrors the Data Preparation section of the notebook. Returns the cleaned
    dataframe used by all query functions.
    """
    if path is None:
        path = DATA_PATH

    df = pd.read_excel(path, header=1)
    df_work = df[USEFUL_COLUMNS].copy().rename(columns=COLUMN_RENAMES)

    # keep only central warehouse, stores and outlets
    df_work = df_work[df_work['WAREHOUSE'].isin([1, 2, 5])].copy()

    # HQ has no shelf location, so missing values are unified under 'HQ'
    df_work['LOCATION'] = df_work['LOCATION'].apply(
        lambda x: str(int(x)) if pd.notna(x) and x != 'HQ' else 'HQ'
    )
    df_work['STORE_ID'] = df_work['WAREHOUSE'].astype(str) + '-' + df_work['LOCATION'].astype(str)

    # portuguese decimal comma, and negatives treated as 0 (conservative)
    df_work['QUANTITY'] = df_work['QUANTITY'].astype(str).str.replace(',', '.').astype(float)
    df_work['QUANTITY'] = df_work['QUANTITY'].clip(lower=0)

    # footwear only; missing family values are excluded as unconfirmable
    df_work = df_work[df_work['FAMILY'].str.contains('CALÇADO', case=False, na=False)].copy()
    df_work['FAMILY'] = 'CALÇADO'

    df_work['BRAND'] = df_work['BRAND'].replace(BRAND_FIXES)

    # scope: the 6 selected stores, keeping HQ as the stock source
    df_work = df_work[df_work['STORE_ID'].isin(SELECTED_STORES + [HQ_ID])].copy()

    # drop brands that exist only at HQ and in no selected store
    relevant_brands = df_work[df_work['STORE_ID'] != HQ_ID]['BRAND'].unique()
    df_work = df_work[df_work['BRAND'].isin(relevant_brands)].copy()

    df_work['SIZE_NUMERIC'] = df_work['SIZE'].apply(convert_size_final)
    df_work['STORE_TYPE'] = df_work['WAREHOUSE'].map(WAREHOUSE_TYPE_LABELS)

    return df_work


# --------------------------------------------------------------------------
# Input helpers
# --------------------------------------------------------------------------

def resolve_store_id(store_input):
    """Resolve a display name, STORE_ID or 'HQ' into the STORE_ID used internally."""
    store_input = str(store_input).strip()

    if store_input.upper() == 'HQ':
        return HQ_ID
    if store_input in SELECTED_STORES:
        return store_input

    reverse_names = {name.lower(): store_id for store_id, name in STORE_NAMES.items()}
    match = reverse_names.get(store_input.lower())
    if match:
        return match

    raise ValueError(
        f"Store not recognized: '{store_input}'. "
        f"Options: {list(STORE_NAMES.values())}, 'HQ', or {SELECTED_STORES}"
    )


def normalize_stores(store_input):
    """None means all selected stores (HQ excluded), a string means one, a list means several."""
    if store_input is None:
        return SELECTED_STORES
    if isinstance(store_input, str):
        store_input = [store_input]
    return [resolve_store_id(s) for s in store_input]


def normalize_brands(df, brand_input):
    """None means all brands present in the data, a string means one, a list means several."""
    all_brands = sorted(df['BRAND'].dropna().unique())
    if brand_input is None:
        return all_brands
    if isinstance(brand_input, str):
        brand_input = [brand_input]

    lookup = {b.lower(): b for b in all_brands}
    resolved = []
    for b in brand_input:
        match = lookup.get(str(b).strip().lower())
        if not match:
            raise ValueError(f"Brand not recognized: '{b}'. Options: {all_brands}")
        resolved.append(match)
    return resolved


def resolve_output_path(filename):
    """Send bare filenames to OUTPUT_DIR, but respect any explicit path given."""
    path = Path(filename)
    if path.parent == Path('.'):
        path = OUTPUT_DIR / path.name
    path.parent.mkdir(parents=True, exist_ok=True)
    return str(path)


# --------------------------------------------------------------------------
# 1 - Stock by store, exported to Excel
# --------------------------------------------------------------------------

def export_stock(df, store_input=None, brand_input=None, filename=None):
    """Export store stock as a size grid, one sheet per store."""
    store_ids = normalize_stores(store_input)
    brands = normalize_brands(df, brand_input)
    multi_brand = len(brands) > 1

    if filename is None:
        filename = 'stock_export.xlsx'
    filename = resolve_output_path(filename)

    data = df[
        (df['STORE_ID'].isin(store_ids)) &
        (df['BRAND'].isin(brands)) &
        (df['SIZE_NUMERIC'].notna())
    ].copy()

    if data.empty:
        store_display = [STORE_NAMES.get(s, s) for s in store_ids]
        raise ValueError(f"No stock of {brands} in {store_display}.")

    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        for store_id in store_ids:
            store_data = data[data['STORE_ID'] == store_id]
            if store_data.empty:
                continue

            # brand only becomes a row level when more than one brand is involved
            index_cols = ['BRAND', 'REFERENCE', 'DESCRIPTION'] if multi_brand else ['REFERENCE', 'DESCRIPTION']

            grid = store_data.pivot_table(
                index=index_cols, columns='SIZE_NUMERIC', values='QUANTITY',
                aggfunc='sum', fill_value=0
            )
            grid = grid.loc[(grid.sum(axis=1) > 0), :]
            if grid.empty:
                continue

            # size scale spans everything shown in this store's sheet
            min_size, max_size = store_data['SIZE_NUMERIC'].min(), store_data['SIZE_NUMERIC'].max()
            full_scale = list(np.arange(min_size, max_size + 0.5, 0.5))
            grid = grid.reindex(columns=full_scale, fill_value=0)
            grid.columns = [format_size_label(c) for c in grid.columns]

            sheet_name = STORE_NAMES.get(store_id, store_id)[:31]
            grid.to_excel(writer, sheet_name=sheet_name)

    wb = load_workbook(filename)
    yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Arial', size=10)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0;-0;""'
        for cell in ws[1]:
            cell.font = Font(name='Arial', size=10, bold=True)
            cell.fill = yellow_fill
        ws.freeze_panes = 'D2' if multi_brand else 'C2'
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_len = max((len(str(c.value)) for c in column_cells if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)
    wb.save(filename)

    return filename


# --------------------------------------------------------------------------
# 2 - Find reference
# --------------------------------------------------------------------------

def find_reference(df, reference_input, store_input=None, size_input=None):
    """Find where a reference exists, matching by prefix.

    Prefix matching is used because the last character of a reference is often an
    internal-system suffix that is not always known upfront.
    """
    ref = str(reference_input).strip().upper()

    matches = df[
        (df['REFERENCE'].astype(str).str.upper().str.startswith(ref)) &
        (df['QUANTITY'] > 0)
    ].copy()

    if store_input is not None:
        matches = matches[matches['STORE_ID'].isin(normalize_stores(store_input))]

    if size_input is not None:
        size_num = convert_size_final(str(size_input))
        matches = matches[matches['SIZE_NUMERIC'] == size_num]

    if matches.empty:
        msg = f"No stock found for reference '{reference_input}'"
        if size_input is not None:
            msg += f" in size {size_input}"
        return msg + "."

    lines = [f"Reference '{reference_input}' found in:"]
    for store_id, store_group in matches.groupby('STORE_ID'):
        store_label = STORE_NAMES.get(store_id, store_id)
        for ref_full, ref_group in store_group.groupby('REFERENCE'):
            sizes_qty = ref_group.groupby('SIZE_NUMERIC')['QUANTITY'].sum()
            sizes_str = ", ".join(
                f"{format_size_label(s)} ({int(q)} units)" for s, q in sizes_qty.items() if q > 0
            )
            lines.append(f"  {store_label} — {ref_full}: {sizes_str}")

    return "\n".join(lines)


# --------------------------------------------------------------------------
# 3 - Quantity summary
# --------------------------------------------------------------------------

def get_quantity(df, store_input=None, brand_input=None):
    """Plain quantity totals, with the output shape adapting to what is specified."""
    store_ids = normalize_stores(store_input)
    brands = normalize_brands(df, brand_input)

    data = df[(df['STORE_ID'].isin(store_ids)) & (df['BRAND'].isin(brands))]
    total = int(data['QUANTITY'].sum())

    multi_store = len(store_ids) > 1
    multi_brand = len(brands) > 1

    # one store, one brand -> single number
    if not multi_store and not multi_brand:
        store_label = STORE_NAMES.get(store_ids[0], store_ids[0])
        return f"{brands[0]} at {store_label}: {total} units."

    # one brand, several stores -> breakdown by store
    if multi_store and not multi_brand:
        lines = [f"{brands[0]} by store:"]
        by_store = data.groupby('STORE_ID')['QUANTITY'].sum()
        for store_id in store_ids:
            lines.append(f"  {STORE_NAMES.get(store_id, store_id)}: {int(by_store.get(store_id, 0))} units")
        lines.append(f"Total: {total} units")
        return "\n".join(lines)

    # one store, several brands -> breakdown by brand
    if not multi_store and multi_brand:
        store_label = STORE_NAMES.get(store_ids[0], store_ids[0])
        lines = [f"{store_label} — quantities by brand:"]
        by_brand = data.groupby('BRAND')['QUANTITY'].sum().sort_values(ascending=False)
        for brand, qty in by_brand.items():
            if qty > 0:
                lines.append(f"  {brand}: {int(qty)} units")
        lines.append(f"Total: {total} units")
        return "\n".join(lines)

    # neither specified -> general breakdown by store
    lines = ["Quantities by store:"]
    by_store = data.groupby('STORE_ID')['QUANTITY'].sum()
    for store_id in store_ids:
        lines.append(f"  {STORE_NAMES.get(store_id, store_id)}: {int(by_store.get(store_id, 0))} units")
    lines.append(f"Total: {total} units")
    return "\n".join(lines)


# --------------------------------------------------------------------------
# 4 - Find model by store type
# --------------------------------------------------------------------------

def find_model_by_type(df, reference_input=None, brand_input=None):
    """Show where a model or brand exists, grouped by HQ / store / outlet."""
    data = df[df['QUANTITY'] > 0].copy()

    if reference_input is not None:
        ref = str(reference_input).strip().upper()
        data = data[data['REFERENCE'].astype(str).str.upper().str.startswith(ref)]

    if brand_input is not None:
        data = data[data['BRAND'].isin(normalize_brands(df, brand_input))]

    if data.empty:
        return "No stock found for that request."

    general = reference_input is None and brand_input is None

    lines = []
    for store_type in ['hq', 'store', 'outlet']:
        type_data = data[data['STORE_TYPE'] == store_type]
        if type_data.empty:
            continue

        type_label = STORE_TYPE_DISPLAY[store_type]

        if general:
            lines.append(f"{type_label}: {int(type_data['QUANTITY'].sum())} units")
        else:
            lines.append(f"{type_label}:")
            for store_id, store_group in type_data.groupby('STORE_ID'):
                store_label = STORE_NAMES.get(store_id, store_id)
                lines.append(f"  {store_label}: {int(store_group['QUANTITY'].sum())} units")

    return "\n".join(lines)


# --------------------------------------------------------------------------
# 5 - Replenishment: missing sizes
# --------------------------------------------------------------------------

REPLENISHMENT_COLUMNS = ['BRAND', 'REFERENCE', 'DESCRIPTION', 'ITEM_TYPE',
                         'SIZE', 'HQ_AVAILABLE', 'SEND_QTY']


def find_replenishable_sizes(df, store_input, brand_input=None, reference_input=None,
                             allocated=None, units_per_size=None):
    """Find sizes HQ has but the store does not, for references the store already carries.

    allocated: {(REFERENCE, SIZE): qty} already committed earlier in the session,
    subtracted from HQ availability so the same units are never proposed twice.
    units_per_size: units to send per missing size (None means everything available).
    """
    store_id = resolve_store_id(store_input)
    if store_id == HQ_ID:
        raise ValueError("HQ is the stock source, not a replenishment target.")

    allocated = allocated or {}
    empty = pd.DataFrame(columns=REPLENISHMENT_COLUMNS)

    data = df[df['SIZE_NUMERIC'].notna()].copy()

    if brand_input is not None:
        data = data[data['BRAND'].isin(normalize_brands(df, brand_input))]

    if reference_input is not None:
        ref = str(reference_input).strip().upper()
        data = data[data['REFERENCE'].astype(str).str.upper().str.startswith(ref)]

    store_data = data[data['STORE_ID'] == store_id]
    hq_data = data[data['STORE_ID'] == HQ_ID]

    # only references the store already carries
    store_refs = store_data.loc[store_data['QUANTITY'] > 0, 'REFERENCE'].unique()
    if len(store_refs) == 0:
        return empty

    # reference/size pairs the store already has covered
    store_sizes = set(
        map(tuple, store_data.loc[store_data['QUANTITY'] > 0, ['REFERENCE', 'SIZE_NUMERIC']].values)
    )

    hq_stock = (
        hq_data[hq_data['REFERENCE'].isin(store_refs)]
        .groupby(['BRAND', 'REFERENCE', 'DESCRIPTION', 'SIZE_NUMERIC'], as_index=False)['QUANTITY'].sum()
    )

    hq_stock['SIZE'] = hq_stock['SIZE_NUMERIC'].apply(format_size_label)

    # discount stock already committed to earlier documents in this session
    hq_stock['HQ_AVAILABLE'] = [
        int(q - allocated.get((r, s), 0))
        for r, s, q in zip(hq_stock['REFERENCE'], hq_stock['SIZE'], hq_stock['QUANTITY'])
    ]

    already_in_store = pd.Series(
        [(r, s) in store_sizes for r, s in zip(hq_stock['REFERENCE'], hq_stock['SIZE_NUMERIC'])],
        index=hq_stock.index
    )

    missing = hq_stock[(hq_stock['HQ_AVAILABLE'] > 0) & (~already_in_store)].copy()
    if missing.empty:
        return empty

    # never send more than HQ actually has
    if units_per_size is None:
        missing['SEND_QTY'] = missing['HQ_AVAILABLE']
    else:
        missing['SEND_QTY'] = missing['HQ_AVAILABLE'].clip(upper=int(units_per_size))

    # item type belongs to the reference, mapped in to avoid splitting rows
    item_types = data.drop_duplicates('REFERENCE').set_index('REFERENCE')['ITEM_TYPE']
    missing['ITEM_TYPE'] = missing['REFERENCE'].map(item_types)

    # alphabetical order, built for physical picking at the warehouse
    missing = missing.sort_values(['BRAND', 'DESCRIPTION', 'REFERENCE', 'SIZE_NUMERIC'])

    return missing[REPLENISHMENT_COLUMNS].reset_index(drop=True)


def generate_replenishment_document(df, store_input, brand_input=None, reference_input=None,
                                    allocated=None, units_per_size=None, filename=None):
    """Generate the printable picking document and return the updated allocation record.

    Returns (filename, new_allocated). This is the only moment stock is considered
    committed; queries alone never allocate anything.
    """
    store_id = resolve_store_id(store_input)
    store_label = STORE_NAMES.get(store_id, store_id)

    result = find_replenishable_sizes(df, store_input, brand_input, reference_input,
                                      allocated, units_per_size)
    if result.empty:
        raise ValueError(f"Nothing to replenish for {store_label} with these filters.")

    if filename is None:
        filename = f"replenishment_{store_id}.xlsx"
    filename = resolve_output_path(filename)

    # HQ_AVAILABLE stays out of the document: only the quantity to pick is printed
    sheet = result[['BRAND', 'REFERENCE', 'DESCRIPTION', 'ITEM_TYPE', 'SIZE', 'SEND_QTY']].rename(
        columns={'SEND_QTY': 'QTY'}
    )
    sheet['PICKED'] = ''

    # leave 3 rows free at the top for the header block
    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        sheet.to_excel(writer, sheet_name=store_label[:31], index=False, startrow=3)

    wb = load_workbook(filename)
    ws = wb.active
    n_cols = sheet.shape[1]
    last_col = get_column_letter(n_cols)
    last_row = 4 + len(sheet)

    ws['A1'] = f"REPLENISHMENT — {store_label.upper()}"
    ws['A1'].font = Font(name='Arial', size=14, bold=True)
    ws.merge_cells(f"A1:{last_col}1")

    info = []
    if brand_input:
        info.append(f"Brand: {brand_input}")
    if reference_input:
        info.append(f"Reference: {reference_input}")
    info.append(f"Lines: {len(sheet)}")
    info.append(f"Total units: {int(result['SEND_QTY'].sum())}")
    info.append(datetime.now().strftime('%d/%m/%Y %H:%M'))

    ws['A2'] = "  |  ".join(info)
    ws['A2'].font = Font(name='Arial', size=9, italic=True)
    ws.merge_cells(f"A2:{last_col}2")

    yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    thin = Side(style='thin', color='999999')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[4]:
        cell.font = Font(name='Arial', size=10, bold=True)
        cell.fill = yellow_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center')

    for row in ws.iter_rows(min_row=5, max_row=last_row, max_col=n_cols):
        for cell in row:
            cell.font = Font(name='Arial', size=10)
            cell.border = border
        row[4].alignment = Alignment(horizontal='center')  # size
        row[5].alignment = Alignment(horizontal='center')  # qty

    for col_idx, column_cells in enumerate(ws.columns, start=1):
        values = [c.value for c in column_cells if c.value is not None and c.row >= 4]
        max_len = max((len(str(v)) for v in values), default=8)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 45)
    ws.column_dimensions[last_col].width = 10  # empty picked column

    # print setup: one page wide, header repeated on every page
    ws.freeze_panes = 'A5'
    ws.print_area = f"A1:{last_col}{last_row}"
    ws.print_title_rows = '4:4'
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.right.text = "Page &P of &N"

    wb.save(filename)

    # commit the stock: this is the only moment it leaves HQ
    new_allocated = dict(allocated or {})
    for ref, size, qty in zip(result['REFERENCE'], result['SIZE'], result['SEND_QTY']):
        new_allocated[(ref, size)] = new_allocated.get((ref, size), 0) + int(qty)

    return filename, new_allocated
